from __future__ import annotations

import io
from typing import Any

import openpyxl


def read_headers(content: bytes) -> list[str]:
    """Возвращает заголовки первой строки xlsx-файла."""
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = wb.active
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return [str(c).strip() if c is not None else "" for c in first_row]
    finally:
        wb.close()


def read_rows(content: bytes) -> list[dict[str, Any]]:
    """Возвращает строки, начиная со второй, как список словарей {заголовок: значение}."""
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        raw_headers = next(rows_iter, ())
        headers = [str(c).strip() if c is not None else "" for c in raw_headers]

        result: list[dict[str, Any]] = []
        for raw_row in rows_iter:
            if raw_row is None or all(v is None for v in raw_row):
                continue
            row = {headers[i]: raw_row[i] for i in range(min(len(headers), len(raw_row)))}
            result.append(row)
        return result
    finally:
        wb.close()
